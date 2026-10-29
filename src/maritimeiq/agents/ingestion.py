"""Ingestion agent: multi-format sources → Bronze Delta with lineage."""
from __future__ import annotations

import hashlib
import re
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from pyspark.sql import DataFrame, Row, SparkSession
from pyspark.sql import functions as F
from pyspark.sql.types import StringType, StructField, StructType

from maritimeiq.mcp.tools import MCPGateway
from maritimeiq.utils.logging_feed import FEED
from maritimeiq.utils.paths import SOURCES, ensure_lakehouse_dirs

LINEAGE = ["_source_file", "_ingested_at", "_row_hash", "_source_system"]


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _sanitize_columns(df: DataFrame) -> DataFrame:
    """Delta / Spark reject spaces and special chars (esp. '.') in column names."""
    used: set[str] = set()
    for c in df.columns:
        leading_us = c.startswith("_")
        safe = re.sub(r"[^A-Za-z0-9_]+", "_", c.strip())
        safe = re.sub(r"_+", "_", safe).strip("_")
        if leading_us:
            safe = f"_{safe}" if not safe.startswith("_") else safe
        if not safe or safe == "_":
            safe = "_col"
        if safe.lstrip("_") and safe.lstrip("_")[0].isdigit():
            safe = f"_c_{safe.lstrip('_')}" if leading_us else f"c_{safe}"
        base = safe
        i = 2
        while safe in used:
            safe = f"{base}_{i}"
            i += 1
        used.add(safe)
        if safe != c:
            df = df.withColumnRenamed(c, safe)
    return df


def _row_hash(values: Iterable[Any]) -> str:
    payload = "|".join("" if v is None else str(v) for v in values)
    return hashlib.sha256(payload.encode("utf-8", errors="ignore")).hexdigest()[:32]


def _add_lineage(df: DataFrame, source_file: str, source_system: str) -> DataFrame:
    cols = df.columns
    # hash over business columns
    concat_cols = F.concat_ws("|", *[F.coalesce(F.col(c).cast("string"), F.lit("")) for c in cols])
    return (
        df.withColumn("_source_file", F.lit(source_file))
        .withColumn("_ingested_at", F.lit(_now()))
        .withColumn("_row_hash", F.sha2(concat_cols, 256))
        .withColumn("_source_system", F.lit(source_system))
    )


class IngestionAgent:
    def __init__(self, spark: SparkSession, gateway: MCPGateway) -> None:
        self.spark = spark
        self.gateway = gateway
        ensure_lakehouse_dirs()

    def run(self, state: dict[str, Any] | None = None) -> dict[str, Any]:
        state = dict(state or {})
        FEED.emit("ingestion", "Starting Bronze ingestion")
        results: dict[str, Any] = {}
        results["ais_positions"] = self.ingest_ais_positions()
        results["ais_live"] = self.ingest_ais_live()
        results["port_agent_calls"] = self.ingest_port_agent_calls()
        results["carrier_schedules"] = self.ingest_carrier_schedules()
        results["bills_of_lading"] = self.ingest_bills_of_lading()
        results["container_events"] = self.ingest_container_events()
        results["crew_manifests"] = self.ingest_crew_manifests()
        results["customers"] = self.ingest_customers()
        results["reference"] = self.ingest_reference()
        state["bronze_tables"] = results
        state["phase"] = "bronze_done"
        FEED.emit("ingestion", "Bronze ingestion complete", results={k: v.get("rows") for k, v in results.items()})
        return state

    def _write(self, df: DataFrame, table: str) -> dict[str, Any]:
        df = _sanitize_columns(df)
        count = df.count()
        out = self.gateway.lakehouse.write_delta(df, f"bronze.{table}", mode="overwrite", agent="ingestion")
        out["rows"] = count
        return out

    def ingest_ais_positions(self) -> dict[str, Any]:
        path = SOURCES / "ais_positions"
        # Hackathon mixes AIS v1 (through 2022-06) and early v2 headers (2022-07+)
        def _header(p: Path) -> str:
            with p.open(encoding="utf-8", errors="ignore") as fh:
                return fh.readline().strip()

        v1_files = [str(p) for p in sorted(path.glob("positions_*.csv")) if _header(p).startswith("MMSI,")]
        v2_files = [str(p) for p in sorted(path.glob("positions_*.csv")) if _header(p).startswith("mmsi,")]
        frames: list[DataFrame] = []
        if v1_files:
            v1 = self.spark.read.option("header", True).csv(v1_files)
            v1 = v1.select(
                F.col("MMSI").alias("MMSI"),
                F.col("IMO").alias("IMO"),
                F.col("VesselName").alias("VesselName"),
                F.col("CallSign").alias("CallSign"),
                F.col("BaseDateTime").alias("BaseDateTime"),
                F.col("LAT").alias("LAT"),
                F.col("LON").alias("LON"),
                F.col("SOG").alias("SOG"),
                F.col("COG").alias("COG"),
                F.col("NavStatus").alias("NavStatus"),
                F.col("Destination").alias("Destination"),
                F.col("ETA").alias("ETA"),
                F.col("Draft").alias("Draft"),
                F.lit("v1").alias("schema_version"),
                F.input_file_name().alias("_src"),
            )
            frames.append(v1)
        if v2_files:
            v2 = self.spark.read.option("header", True).csv(v2_files)
            # Convert epoch ms to ISO-ish string for a unified Bronze column
            v2 = v2.select(
                F.col("mmsi").alias("MMSI"),
                F.col("imo_number").cast("string").alias("IMO"),
                F.col("ship_name").alias("VesselName"),
                F.col("call_sign").alias("CallSign"),
                F.from_unixtime((F.col("ts_utc_ms").cast("long") / 1000).cast("long")).alias("BaseDateTime"),
                F.col("latitude").alias("LAT"),
                F.col("longitude").alias("LON"),
                F.col("speed_knots").alias("SOG"),
                F.col("course").alias("COG"),
                F.col("nav_status").alias("NavStatus"),
                F.col("dest").alias("Destination"),
                F.col("eta_utc").alias("ETA"),
                F.col("draught_m").alias("Draft"),
                F.lit("v2").alias("schema_version"),
                F.input_file_name().alias("_src"),
            )
            frames.append(v2)
        df = frames[0]
        for other in frames[1:]:
            df = df.unionByName(other, allowMissingColumns=True)
        business = [c for c in df.columns if not c.startswith("_")]
        df = (
            df.withColumn("_source_file", F.col("_src"))
            .withColumn("_ingested_at", F.lit(_now()))
            .withColumn(
                "_row_hash",
                F.sha2(
                    F.concat_ws("|", *[F.coalesce(F.col(c).cast("string"), F.lit("")) for c in business]),
                    256,
                ),
            )
            .withColumn("_source_system", F.lit("ais_positions"))
            .drop("_src")
        )
        FEED.emit("ingestion", f"ais_positions loaded v1={len(v1_files)} v2={len(v2_files)}")
        return self._write(df, "ais_positions")

    def ingest_ais_live(self) -> dict[str, Any]:
        files = sorted(str(p) for p in (SOURCES / "ais_live").glob("*.csv"))
        if not files:
            return {"ok": False, "rows": 0}
        df = self.spark.read.option("header", True).csv(files)
        df = _add_lineage(df, ",".join(files), "ais_live")
        return self._write(df, "ais_live")

    def ingest_port_agent_calls(self) -> dict[str, Any]:
        root = SOURCES / "port_agent_calls"
        # Read all jsonl; spark json reader handles json lines
        df = self.spark.read.option("multiLine", False).json(str(root / "agency=*/*.jsonl"))
        # Keep agency from path
        df = df.withColumn("_src", F.input_file_name())
        df = (
            df.withColumn(
                "agency",
                F.regexp_extract(F.col("_src"), r"agency=([^/]+)", 1),
            )
            .withColumn("_source_file", F.col("_src"))
            .withColumn("_ingested_at", F.lit(_now()))
            .withColumn("_row_hash", F.sha2(F.to_json(F.struct(*[c for c in df.columns if c != "_src"])), 256))
            .withColumn("_source_system", F.lit("port_agent_calls"))
            .drop("_src")
        )
        # Cast complex to JSON strings for a flatter Bronze table usable in SQL demos
        flat_cols = []
        select_exprs = []
        for c, t in df.dtypes:
            if t.startswith("struct") or t.startswith("array") or t == "map":
                select_exprs.append(F.to_json(F.col(c)).alias(c))
            else:
                select_exprs.append(F.col(c))
            flat_cols.append(c)
        df2 = df.select(*select_exprs)
        FEED.emit("ingestion", "port_agent_calls loaded")
        return self._write(df2, "port_agent_calls")

    def ingest_carrier_schedules(self) -> dict[str, Any]:
        path = SOURCES / "carrier_schedules"
        raw = self.spark.read.text(str(path / "*.txt"))
        raw = raw.withColumn("_src", F.input_file_name())
        # Keep D records only for the main bronze table
        d = raw.where(F.col("value").startswith("D"))
        parsed = (
            d.withColumn("rec_type", F.substring("value", 1, 1))
            .withColumn("service", F.trim(F.substring("value", 2, 6)))
            .withColumn("voyage", F.trim(F.substring("value", 8, 10)))
            .withColumn("vessel_name", F.trim(F.substring("value", 18, 25)))
            .withColumn("imo", F.trim(F.substring("value", 43, 7)))
            .withColumn("port_locode", F.trim(F.substring("value", 50, 5)))
            .withColumn("terminal", F.trim(F.substring("value", 55, 6)))
            .withColumn("eta_raw", F.trim(F.substring("value", 61, 12)))
            .withColumn("etd_raw", F.trim(F.substring("value", 73, 12)))
            .withColumn("status", F.trim(F.substring("value", 85, 1)))
            .withColumn("_source_file", F.col("_src"))
            .withColumn("_ingested_at", F.lit(_now()))
            .withColumn("_source_system", F.lit("carrier_schedules"))
            .drop("value", "_src")
        )
        parsed = parsed.withColumn(
            "_row_hash",
            F.sha2(
                F.concat_ws(
                    "|",
                    *[F.coalesce(F.col(c).cast("string"), F.lit("")) for c in parsed.columns if not c.startswith("_")],
                ),
                256,
            ),
        )
        # SCAC from filename
        parsed = parsed.withColumn(
            "scac",
            F.regexp_extract(F.col("_source_file"), r"/([A-Z]{4})_\d{4}Q", 1),
        )
        FEED.emit("ingestion", "carrier_schedules loaded")
        return self._write(parsed, "carrier_schedules")

    def ingest_bills_of_lading(self) -> dict[str, Any]:
        """Parse B/L XML via ElementTree (avoids spark-xml package download)."""
        rows: list[dict[str, Any]] = []
        root = SOURCES / "bills_of_lading"
        for fp in sorted(root.glob("bl_*.xml")):
            try:
                tree = ET.parse(fp)
            except ET.ParseError:
                continue
            xml_root = tree.getroot()
            # Handle with or without namespace
            bills = list(xml_root.iter())
            # If root itself is BillOfLading or file has multiple
            candidates = []
            if xml_root.tag.endswith("BillOfLading"):
                candidates = [xml_root]
            else:
                candidates = [e for e in xml_root.iter() if e.tag.endswith("BillOfLading")]
            if not candidates:
                # sometimes root wraps many
                candidates = [xml_root] if "BLNumber" in {c.tag.split('}')[-1] for c in xml_root} else []
            # Fallback: treat each top-level child
            if not candidates:
                candidates = list(xml_root)

            def text(el: ET.Element | None, *path: str) -> str | None:
                if el is None:
                    return None
                cur = el
                for p in path:
                    found = None
                    for ch in cur:
                        if ch.tag.endswith(p):
                            found = ch
                            break
                    if found is None:
                        return None
                    cur = found
                return (cur.text or "").strip() if cur is not None else None

            for bl in candidates:
                if not any(ch.tag.endswith("BLNumber") for ch in bl):
                    continue
                containers = []
                for cnode in bl.iter():
                    if cnode.tag.endswith("Container") and any(x.tag.endswith("Number") for x in cnode):
                        containers.append(
                            {
                                "number": text(cnode, "Number"),
                                "type": text(cnode, "Type"),
                                "seal": text(cnode, "Seal"),
                                "gross_weight": text(cnode, "GrossWeight"),
                            }
                        )
                if not containers:
                    containers = [{"number": None, "type": None, "seal": None, "gross_weight": None}]
                for cont in containers:
                    row = {
                        "BLNumber": text(bl, "BLNumber"),
                        "BookingRef": text(bl, "BookingRef"),
                        "version": bl.attrib.get("version"),
                        "Carrier": text(bl, "Carrier"),
                        "IssueDate": text(bl, "IssueDate"),
                        "ShipperName": text(bl, "Shipper", "Name"),
                        "ShipperContact": text(bl, "Shipper", "Contact"),
                        "ShipperPhone": text(bl, "Shipper", "Phone"),
                        "ShipperEmail": text(bl, "Shipper", "Email"),
                        "ShipperTaxId": text(bl, "Shipper", "TaxId"),
                        "ConsigneeName": text(bl, "Consignee", "Name"),
                        "ConsigneeContact": text(bl, "Consignee", "Contact"),
                        "ConsigneePhone": text(bl, "Consignee", "Phone"),
                        "ConsigneeEmail": text(bl, "Consignee", "Email"),
                        "ConsigneeTaxId": text(bl, "Consignee", "TaxId"),
                        "Vessel": text(bl, "Vessel"),
                        "Voyage": text(bl, "Voyage"),
                        "PortOfLoading": text(bl, "PortOfLoading"),
                        "PortOfDischarge": text(bl, "PortOfDischarge"),
                        "ContainerNumber": cont["number"],
                        "ContainerType": cont["type"],
                        "Seal": cont["seal"],
                        "GrossWeight": cont["gross_weight"],
                        "_source_file": str(fp),
                        "_ingested_at": _now(),
                        "_source_system": "bills_of_lading",
                    }
                    row["_row_hash"] = _row_hash(row.values())
                    rows.append(row)
        if not rows:
            FEED.emit("ingestion", "bills_of_lading empty", level="warn")
            schema = StructType([StructField("BLNumber", StringType(), True)])
            df = self.spark.createDataFrame([], schema)
        else:
            df = self.spark.createDataFrame(rows)
        FEED.emit("ingestion", f"bills_of_lading rows={len(rows)}")
        return self._write(df, "bills_of_lading")

    def ingest_container_events(self) -> dict[str, Any]:
        root = SOURCES / "container_events"
        frames: list[DataFrame] = []

        # DCSA CSV
        dcsa_carriers = ["MAEU", "CMDU", "HLCU", "ZIMU"]
        for scac in dcsa_carriers:
            p = root / scac
            if not p.exists():
                continue
            df = self.spark.read.option("header", True).csv(str(p / "*.csv"))
            df = (
                df.withColumn("feed_format", F.lit("dcsa"))
                .withColumn("carrier_scac", F.lit(scac))
                .withColumn("_src", F.input_file_name())
            )
            business = [c for c in df.columns if not c.startswith("_")]
            df = (
                df.withColumn("_source_file", F.col("_src"))
                .withColumn("_ingested_at", F.lit(_now()))
                .withColumn(
                    "_row_hash",
                    F.sha2(F.concat_ws("|", *[F.coalesce(F.col(c).cast("string"), F.lit("")) for c in business]), 256),
                )
                .withColumn("_source_system", F.lit("container_events"))
                .drop("_src")
            )
            frames.append(df)

        # EDI semicolon
        for scac in ["MSCU", "ONEY"]:
            p = root / scac
            if not p.exists():
                continue
            df = self.spark.read.option("header", True).option("sep", ";").csv(str(p / "*.txt"))
            df = df.withColumn("feed_format", F.lit("edi_semicolon")).withColumn("carrier_scac", F.lit(scac))
            df = df.withColumn("_src", F.input_file_name())
            business = [c for c in df.columns if not c.startswith("_")]
            df = (
                df.withColumn("_source_file", F.col("_src"))
                .withColumn("_ingested_at", F.lit(_now()))
                .withColumn(
                    "_row_hash",
                    F.sha2(F.concat_ws("|", *[F.coalesce(F.col(c).cast("string"), F.lit("")) for c in business]), 256),
                )
                .withColumn("_source_system", F.lit("container_events"))
                .drop("_src")
            )
            frames.append(df)

        # Pipe legacy — apply schema (no header pre-2023)
        pipe_schema = [
            "container",
            "evt",
            "mt_fl",
            "loc",
            "evt_date",
            "evt_time",
            "tz",
            "vsl_callsign",
            "voyage",
            "bkg_ref",
            "seq",
        ]
        for scac in ["COSU", "EGLV", "YMLU"]:
            p = root / scac
            if not p.exists():
                continue
            # read without header; if first line looks like header, spark still fine as data
            df = (
                self.spark.read.option("header", False)
                .option("sep", "|")
                .csv(str(p / "*.dat"))
                .toDF(*pipe_schema[:])
            )
            # Drop accidental header rows
            df = df.where(F.lower(F.col("container")) != F.lit("container"))
            df = df.withColumn("feed_format", F.lit("pipe_legacy")).withColumn("carrier_scac", F.lit(scac))
            df = df.withColumn("_src", F.input_file_name())
            business = [c for c in df.columns if not c.startswith("_")]
            df = (
                df.withColumn("_source_file", F.col("_src"))
                .withColumn("_ingested_at", F.lit(_now()))
                .withColumn(
                    "_row_hash",
                    F.sha2(F.concat_ws("|", *[F.coalesce(F.col(c).cast("string"), F.lit("")) for c in business]), 256),
                )
                .withColumn("_source_system", F.lit("container_events"))
                .drop("_src")
            )
            frames.append(df)

        # Union by name
        if not frames:
            return {"ok": False, "rows": 0}
        base = frames[0]
        for other in frames[1:]:
            base = base.unionByName(other, allowMissingColumns=True)
        FEED.emit("ingestion", "container_events loaded")
        return self._write(base, "container_events")

    def ingest_crew_manifests(self) -> dict[str, Any]:
        rows: list[dict[str, Any]] = []
        root = SOURCES / "crew_manifests"
        for fp in sorted(root.glob("crew_*.xlsx")):
            wb = load_workbook(fp, read_only=True, data_only=True)
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                data = list(ws.iter_rows(values_only=True))
                if not data:
                    continue
                # Find header row containing Passport or Surname / Family Name
                header_idx = None
                for i, row in enumerate(data[:15]):
                    cells = [str(c).strip().lower() if c is not None else "" for c in row]
                    joined = " ".join(cells)
                    if "passport" in joined or "surname" in joined or "family name" in joined:
                        header_idx = i
                        break
                if header_idx is None:
                    continue
                headers = [str(c).strip() if c is not None else f"col{j}" for j, c in enumerate(data[header_idx])]
                # Header block is "Label:" / value pairs (ship, IMO, port of arrival, date).
                meta: dict[str, str] = {}
                for row in data[:header_idx]:
                    vals = [c for c in row if c is not None and str(c).strip()]
                    for j, v in enumerate(vals[:-1]):
                        label = str(v).strip().lower()
                        if label.endswith(":"):
                            meta[label] = str(vals[j + 1]).strip()
                ship_name = next((v for k, v in meta.items() if "ship" in k or "vessel" in k), None)
                imo_txt = next((v for k, v in meta.items() if "imo" in k), "")
                imo = next((t for t in re.findall(r"\d{7}", imo_txt)), None)
                port = next((v for k, v in meta.items() if "port" in k), None)
                arrival = next((v for k, v in meta.items() if k.startswith("date")), None)
                for row in data[header_idx + 1 :]:
                    if all(c is None or str(c).strip() == "" for c in row):
                        continue
                    rec = {headers[j]: (None if j >= len(row) else row[j]) for j in range(len(headers))}
                    # normalize keys
                    norm = {str(k): (str(v) if v is not None else None) for k, v in rec.items()}
                    norm["vessel_name_meta"] = ship_name
                    norm["imo_meta"] = imo
                    norm["port_meta"] = port
                    norm["arrival_date_meta"] = arrival
                    norm["sheet"] = sheet
                    norm["_source_file"] = str(fp)
                    norm["_ingested_at"] = _now()
                    norm["_source_system"] = "crew_manifests"
                    norm["_row_hash"] = _row_hash(norm.values())
                    rows.append(norm)
            wb.close()
        if not rows:
            df = self.spark.createDataFrame([], StructType([StructField("Passport", StringType(), True)]))
        else:
            # unify keys
            keys = sorted({k for r in rows for k in r})
            unified = [{k: r.get(k) for k in keys} for r in rows]
            df = self.spark.createDataFrame([Row(**u) for u in unified])
        # Register PII tags
        pii_cols = [c for c in df.columns if any(x in c.lower() for x in ["passport", "dob", "birth", "surname", "given", "family", "first name", "seaman", "cdc"])]
        from maritimeiq.governance.guard import GUARD

        GUARD.register_pii("bronze.crew_manifests", pii_cols)
        FEED.emit("ingestion", f"crew_manifests rows={len(rows)} pii_cols={pii_cols}")
        return self._write(df, "crew_manifests")

    def ingest_customers(self) -> dict[str, Any]:
        path = SOURCES / "customers" / "crm_accounts_export.csv"
        df = self.spark.read.option("header", True).csv(str(path))
        df = _add_lineage(df, str(path), "customers")
        from maritimeiq.governance.guard import GUARD

        pii = [c for c in df.columns if any(x in c.lower() for x in ["contact", "email", "phone", "street", "tax", "owner"])]
        GUARD.register_pii("bronze.customers", pii)
        FEED.emit("ingestion", "customers loaded")
        return self._write(df, "customers")

    def ingest_reference(self) -> dict[str, Any]:
        out = {}
        for name in ["unlocode_code_list", "world_port_index", "carriers", "carrier_alliances"]:
            path = SOURCES / "reference" / f"{name}.csv"
            if not path.exists():
                continue
            df = self.spark.read.option("header", True).csv(str(path))
            df = _add_lineage(df, str(path), "reference")
            out[name] = self._write(df, f"ref_{name}")
        return {"ok": True, "tables": out}
